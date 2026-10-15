"""Inventory a product folder and turn every document into API content blocks.

Rules:
- PDFs travel as PDF documents (the model reads scans visually).
- Heavy PDFs (artwork, labels) are rasterised page by page to JPEG so the request stays
  under the API limit; text layer is attached alongside so nothing is lost.
- DOCX / XLSX / TXT are converted to plain text locally.
- Images go as images.
- Anything else is listed by name only.
Nothing is modified inside the folder.
"""
from __future__ import annotations

import base64
import io
import os
from dataclasses import dataclass, field
from pathlib import Path

IGNORED_DIRS = {"_invimagil", "_referencia", ".git", "__MACOSX"}
IGNORED_FILES = {".DS_Store", "Thumbs.db", "desktop.ini"}
# outputs of this tool or of earlier manual runs: never fed back as source documents
IGNORED_PREFIXES = ("MAESTRO NSO", "log-nso", "planilla", "informe-nso", "RESUMEN-")
PDF_INLINE_MAX_MB = 3.0          # above this a PDF is rasterised instead of sent as PDF
RASTER_DPI = 110
JPEG_QUALITY = 72
REQUEST_BUDGET_MB = 24.0         # stay well under the 32 MB request limit
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".gif", ".webp"}


@dataclass
class Doc:
    path: Path
    rel: str
    kind: str                      # pdf | pdf_raster | docx | xlsx | text | image | other
    size_mb: float
    pages: int = 0
    blocks: list[dict] = field(default_factory=list)
    note: str = ""


def list_files(folder: Path) -> list[Path]:
    out: list[Path] = []
    for root, dirs, files in os.walk(folder):
        dirs[:] = sorted(d for d in dirs if d not in IGNORED_DIRS and not d.startswith("."))
        for f in sorted(files):
            if f in IGNORED_FILES or f.startswith("~$") or f.startswith(".") or f.startswith(IGNORED_PREFIXES):
                continue
            out.append(Path(root) / f)
    return out


def _b64(data: bytes) -> str:
    return base64.standard_b64encode(data).decode("ascii")


def _pdf_blocks(path: Path, raster: bool) -> tuple[list[dict], int, str]:
    import pymupdf

    doc = pymupdf.open(path)
    pages = doc.page_count
    blocks: list[dict] = []
    if not raster:
        blocks.append(
            {
                "type": "document",
                "source": {"type": "base64", "media_type": "application/pdf", "data": _b64(path.read_bytes())},
                "title": path.name,
            }
        )
        return blocks, pages, ""
    text_parts = []
    for i, page in enumerate(doc):
        pix = page.get_pixmap(dpi=RASTER_DPI, colorspace=pymupdf.csRGB, alpha=False)
        jpg = pix.tobytes("jpeg", jpg_quality=JPEG_QUALITY)
        blocks.append({"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": _b64(jpg)}})
        t = page.get_text().strip()
        if t:
            text_parts.append(f"--- página {i + 1} ---\n{t}")
    if text_parts:
        blocks.append({"type": "text", "text": f"Capa de texto de «{path.name}»:\n" + "\n".join(text_parts)})
    return blocks, pages, f"rasterizado a {RASTER_DPI} dpi ({pages} pág.)"


def _docx_text(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for ti, table in enumerate(d.tables, 1):
        parts.append(f"[Tabla {ti}]")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def _xlsx_text(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"[Hoja {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                parts.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(parts)


def prepare_docs(folder: Path, progress=lambda msg: None) -> list[Doc]:
    docs: list[Doc] = []
    for p in list_files(folder):
        rel = str(p.relative_to(folder))
        size_mb = p.stat().st_size / 1e6
        ext = p.suffix.lower()
        d = Doc(path=p, rel=rel, kind="other", size_mb=round(size_mb, 2))
        try:
            if ext == ".pdf":
                raster = size_mb > PDF_INLINE_MAX_MB
                d.kind = "pdf_raster" if raster else "pdf"
                d.blocks, d.pages, d.note = _pdf_blocks(p, raster)
            elif ext == ".docx":
                d.kind = "docx"
                d.blocks = [{"type": "text", "text": f"Contenido de «{rel}» (Word):\n{_docx_text(p)}"}]
            elif ext in {".xlsx", ".xlsm"}:
                d.kind = "xlsx"
                d.blocks = [{"type": "text", "text": f"Contenido de «{rel}» (Excel):\n{_xlsx_text(p)}"}]
            elif ext in {".txt", ".md", ".csv"}:
                d.kind = "text"
                d.blocks = [{"type": "text", "text": f"Contenido de «{rel}»:\n{p.read_text(encoding='utf-8', errors='replace')}"}]
            elif ext in IMAGE_EXT:
                d.kind = "image"
                mt = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "gif": "image/gif", "webp": "image/webp"}[ext[1:]]
                d.blocks = [{"type": "image", "source": {"type": "base64", "media_type": mt, "data": _b64(p.read_bytes())}}]
            else:
                d.note = "formato no leído; solo se registra el nombre"
        except Exception as exc:  # noqa: BLE001 - report, never abort the whole folder
            d.kind = "other"
            d.blocks = []
            d.note = f"no se pudo leer: {exc}"
        progress(f"  {d.rel}  ({d.size_mb} MB, {d.kind}{', ' + d.note if d.note else ''})")
        docs.append(d)
    _fit_budget(docs, progress)
    return docs


def payload_mb(docs: list[Doc]) -> float:
    total = 0
    for d in docs:
        for b in d.blocks:
            if b["type"] in ("document", "image"):
                total += len(b["source"]["data"])
            else:
                total += len(b["text"].encode("utf-8"))
    return total / 1e6


def _fit_budget(docs: list[Doc], progress) -> None:
    """If the payload is still too big, rasterise the largest inline PDFs."""
    while payload_mb(docs) > REQUEST_BUDGET_MB:
        candidates = [d for d in docs if d.kind == "pdf"]
        if not candidates:
            break
        big = max(candidates, key=lambda d: d.size_mb)
        big.blocks, big.pages, big.note = _pdf_blocks(big.path, raster=True)
        big.kind = "pdf_raster"
        progress(f"  {big.rel}: convertido a imágenes para ajustar el tamaño del envío")


def build_user_content(docs: list[Doc], intro: str) -> list[dict]:
    content: list[dict] = [{"type": "text", "text": intro}]
    for i, d in enumerate(docs, 1):
        header = f"\n=== Archivo {i}/{len(docs)}: «{d.rel}» · {d.size_mb} MB"
        if d.pages:
            header += f" · {d.pages} página(s)"
        if d.note:
            header += f" · {d.note}"
        content.append({"type": "text", "text": header})
        content.extend(d.blocks)
    return content


def inventory_table(docs: list[Doc]) -> str:
    lines = ["| Archivo | MB | Páginas | Cómo se leyó |", "|---|---:|---:|---|"]
    for d in docs:
        lines.append(f"| {d.rel} | {d.size_mb} | {d.pages or ''} | {d.kind}{' · ' + d.note if d.note else ''} |")
    return "\n".join(lines)

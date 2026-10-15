"""planilla.json -> MAESTRO NSO Excel (one sheet per InvimÁgil wizard step, copy-paste ready)."""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .catalogos import Catalog

FONT = "Arial"
GREY = "D9D9D9"
BLUE = "DDEBF7"
LIGHT = "F2F2F2"
YELLOW = "FFF2CC"
RED = "C00000"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STEP_COLS = ["#", "Campo del formulario", "Valor para copiar", "Tipo de campo", "Código catálogo", "Estado", "Nota / fuente"]
STEP_WIDTHS = [9, 36, 64, 14, 16, 16, 60]


def _font(bold=False, color="000000", size=10, italic=False) -> Font:
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def _fill(hex_: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_)


class SheetWriter:
    def __init__(self, ws, widths: list[int]):
        self.ws = ws
        self.row = 1
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def title(self, text: str, subtitle: str = ""):
        c = self.ws.cell(self.row, 1, text)
        c.font = _font(bold=True, size=13)
        self.row += 1
        if subtitle:
            c = self.ws.cell(self.row, 1, subtitle)
            c.font = _font(italic=True, color="595959")
            self.row += 1
        self.row += 1

    def header(self, cols: list[str]):
        for i, h in enumerate(cols, 1):
            c = self.ws.cell(self.row, i, h)
            c.font = _font(bold=True)
            c.fill = _fill(GREY)
            c.border = BORDER
            c.alignment = Alignment(vertical="center", wrap_text=True)
        self.ws.freeze_panes = self.ws.cell(self.row + 1, 1)
        self.row += 1

    def section(self, text: str, ncols: int = 7):
        for i in range(1, ncols + 1):
            c = self.ws.cell(self.row, i)
            c.fill = _fill(BLUE)
        self.ws.cell(self.row, 1, text).font = _font(bold=True)
        self.row += 1

    def subsection(self, text: str, ncols: int = 7):
        for i in range(1, ncols + 1):
            self.ws.cell(self.row, i).fill = _fill(LIGHT)
        self.ws.cell(self.row, 2, text).font = _font(bold=True)
        self.row += 1

    def blank(self):
        self.row += 1

    def line(self, values: list, wrap_cols: tuple[int, ...] = (3, 7), pending: bool = False, input_col: int | None = None):
        for i, v in enumerate(values, 1):
            c = self.ws.cell(self.row, i, v)
            c.font = _font()
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=(i in wrap_cols))
        if pending:
            cell = self.ws.cell(self.row, 6)
            cell.font = _font(bold=True, color=RED)
        if input_col:
            self.ws.cell(self.row, input_col).fill = _fill(YELLOW)
        self.row += 1


def _state(node: dict | None) -> tuple[str, bool]:
    if node and node.get("_decision"):
        return "Pendiente", True
    return "Listo", False


def _val(v):
    if v is None or v == "":
        return "PENDIENTE"
    if isinstance(v, bool):
        return "Sí" if v else "No"
    return v


def _cat_cells(node: dict | None) -> tuple[str, str, str, bool, str]:
    """-> (valor, código, estado, pendiente, nota)"""
    node = node or {}
    state, pend = _state(node)
    name = node.get("name")
    slug = node.get("slug")
    if not name and not slug:
        return "PENDIENTE", "", "Pendiente", True, node.get("_decision") or ""
    return name or "PENDIENTE", slug or "", state, pend, node.get("_decision") or node.get("_note") or ""


def _field(sw: SheetWriter, num: str, label: str, node: dict | None, key: str, kind="Texto"):
    node = node or {}
    missing = node.get(key) in (None, "")
    note = (node.get("_decision") if missing else None) or node.get("_note") or ""
    sw.line([num, label, _val(node.get(key)), kind, "", "Pendiente" if missing else "Listo", note], pending=missing)


def _catfield(sw: SheetWriter, num: str, label: str, node: dict | None):
    valor, codigo, estado, pend, nota = _cat_cells(node)
    sw.line([num, label, valor, "Lista", codigo, estado, nota], pending=pend)


def _file_mb(folder: Path, name: str) -> str:
    p = folder / name
    if p.exists():
        return f"{p.stat().st_size / 1e6:.2f}"
    for q in folder.rglob("*"):
        if q.name == name:
            return f"{q.stat().st_size / 1e6:.2f}"
    return "no encontrado"


def write_excel(planilla: dict, folder: Path, out_path: Path, cat: Catalog, resueltas: list[dict] | None = None) -> None:
    wb = Workbook()
    meta = planilla.get("_meta", {})
    general = planilla.get("general", {})
    today = dt.date.today().strftime("%d/%m/%Y")

    # --- Portada ---------------------------------------------------------------
    ws = wb.active
    ws.title = "Portada"
    sw = SheetWriter(ws, [30, 100])
    sw.title("Maestro NSO · InvimÁgil", "Ficha maestra del producto para diligenciar el formulario paso a paso (copiar y pegar).")
    sw.section("Identificación del trámite", 2)
    for k, v in [
        ("Producto", meta.get("producto")),
        ("Tipo de trámite", meta.get("tipoTramite")),
        ("Titular de la NSO", meta.get("titular")),
        ("Fabricante", meta.get("fabricanteResumen")),
        ("Fuente de los datos", meta.get("fuente")),
        ("Carpeta", str(folder)),
        ("Fecha de elaboración", f"{today} (generado por la herramienta InvimÁgil Base; revisar antes de usar)"),
    ]:
        sw.ws.cell(sw.row, 1, k).font = _font(bold=True)
        c = sw.ws.cell(sw.row, 2, _val(v))
        c.font = _font()
        c.alignment = Alignment(wrap_text=True, vertical="top")
        sw.row += 1
    sw.blank()
    sw.section("Orden del formulario (una hoja por paso del wizard)", 2)
    for k, v in [
        ("1 Información general", "Responsables, origen, fabricante, datos del producto y presentaciones"),
        ("2 Composición", "Ingredientes INCI en el orden del COCC, con función, listado y porcentaje"),
        ("3 Información de uso", "Modo de uso y advertencias"),
        ("4 Caract. técnicas", "Organolépticas, fisicoquímicas, microbiológicas y vida útil"),
        ("5 Doc. obligatorios", "Archivo literal a adjuntar en cada casilla"),
        ("6 Doc. adicionales", "Archivos adicionales propuestos"),
        ("Decisiones", "Lo que falta decidir. Responda en la columna amarilla y use «Aplicar decisiones»"),
        ("Archivos", "Cómo se clasificó cada archivo de la carpeta"),
    ]:
        sw.ws.cell(sw.row, 1, k).font = _font(bold=True)
        sw.ws.cell(sw.row, 2, v).font = _font()
        sw.row += 1
    sw.blank()
    sw.section("Hallazgos para revisar antes de radicar", 2)
    for h in meta.get("hallazgos") or ["(ninguno)"]:
        c = sw.ws.cell(sw.row, 2, f"• {h}")
        c.font = _font()
        c.alignment = Alignment(wrap_text=True, vertical="top")
        sw.row += 1
    sw.blank()
    sw.section("Convenciones", 2)
    for k, v in [
        ("Listo", "Dato tomado de los documentos y validado contra el catálogo de InvimÁgil."),
        ("Pendiente (rojo)", "Falta una decisión. Está explicada en la hoja Decisiones."),
        ("Código catálogo", "Identificador interno de la opción en la lista de InvimÁgil; sirve para verificar que se eligió la opción correcta."),
    ]:
        sw.ws.cell(sw.row, 1, k).font = _font(bold=True)
        sw.ws.cell(sw.row, 2, v).font = _font()
        sw.row += 1

    # --- 1 Información general -------------------------------------------------
    ws = wb.create_sheet("1 Información general")
    sw = SheetWriter(ws, STEP_WIDTHS)
    sw.title("Paso 1 · Información general", "Responsables, origen, fabricante, producto y presentaciones.")
    sw.header(STEP_COLS)
    resp = general.get("responsibles", {})
    sw.section("1.1 Responsables")
    _field(sw, "1.1.1", "Responsable técnico", resp.get("technical"), "name", "Lista")
    _field(sw, "1.1.2", "Representante legal / apoderado", resp.get("legal"), "name", "Lista")
    sw.blank()
    sw.section("1.2 Origen")
    origin = general.get("origin", {})
    _field(sw, "1.2.1", "Origen del producto", origin, "value", "Lista")
    if origin.get("importer"):
        sw.line(["1.2.2", "Importador (según etiqueta/sticker)", origin["importer"], "Texto", "", "Informativo", ""])
    sw.blank()
    sw.section("1.3 Fabricante")
    manu = general.get("manufacturer", {})
    _field(sw, "1.3.1", "Tipo de fabricante", manu, "type", "Lista")
    for i, t in enumerate(manu.get("thirds") or [], 1):
        co = t.get("company", {})
        sw.subsection(f"Fabricante {i}")
        sw.line([f"1.3.{i}.1", "Razón social", _val(co.get("name")), "Texto / Lista", "", "Listo", co.get("_note") or ""])
        sw.line([f"1.3.{i}.2", "Dirección", _val(co.get("address")), "Texto", "", "Listo" if co.get("address") else "Pendiente", ""], pending=not co.get("address"))
        sw.line([f"1.3.{i}.3", "País", f"{co.get('countryName') or 'PENDIENTE'} ({co.get('countryCode') or '?'})", "Lista", "", "Listo", ""])
        sw.line([f"1.3.{i}.4", "¿Extranjero?", _val(co.get("isForeign")), "Sí/No", "", "Listo", ""])
    sw.blank()
    sw.section("1.4 Producto")
    prod = general.get("product", {})
    _field(sw, "1.4.1", "Nombre del producto", prod, "name")
    _field(sw, "1.4.2", "Denominación genérica", prod, "genericName")
    sw.line(["1.4.3", "Marca(s)", ", ".join(prod.get("brands") or []) or "PENDIENTE", "Texto", "", "Listo" if prod.get("brands") else "Pendiente", ""], pending=not prod.get("brands"))
    sw.line(["1.4.4", "¿Declara grupos cosméticos?", _val(prod.get("declaresCosmeticGroups")), "Sí/No", "", "Listo", ""])
    _catfield(sw, "1.4.5", "Forma cosmética", prod.get("cosmeticForm"))
    _catfield(sw, "1.4.6", "Categoría del producto", prod.get("productCategory"))
    if prod.get("_decision") or prod.get("_note"):
        sw.line(["", "Nota del producto", prod.get("_decision") or prod.get("_note"), "", "", "Pendiente" if prod.get("_decision") else "Informativo", ""], pending=bool(prod.get("_decision")))
    sw.blank()
    sw.section("1.5 Presentaciones comerciales")
    for p in general.get("presentations") or []:
        i = p.get("order")
        unit = (p.get("measureUnit") or {}).get("name") or "?"
        sw.subsection(f"Presentación {i} · {p.get('netContent')} {unit}")
        sw.line([f"1.5.{i}.1", "Contenido neto", p.get("netContent"), "Número", "", "Listo", ""])
        _catfield(sw, f"1.5.{i}.2", "Unidad de medida", p.get("measureUnit"))
        _catfield(sw, f"1.5.{i}.3", "Tipo de envase primario", p.get("primaryContainer"))
        _catfield(sw, f"1.5.{i}.4", "Material del envase primario", p.get("primaryMaterial"))
        _catfield(sw, f"1.5.{i}.5", "Tipo de envase secundario", p.get("secondaryContainer"))
        _catfield(sw, f"1.5.{i}.6", "Material del envase secundario", p.get("secondaryMaterial"))
        sw.line([f"1.5.{i}.7", "Observaciones", p.get("observations") or "", "Texto", "", "Listo", p.get("_decision") or ""], pending=bool(p.get("_decision")))

    # --- 2 Composición ------------------------------------------------------------
    ws = wb.create_sheet("2 Composición")
    cols = ["#", "Nombre INCI", "N° CAS", "Tipo de ingrediente", "Función (catálogo InvimÁgil)", "Código función", "Listado de referencia", "Cantidad", "Unidad", "¿Nanomaterial?", "Función según COCC", "Estado", "Nota / decisión"]
    sw = SheetWriter(ws, [6, 44, 14, 16, 42, 14, 18, 14, 10, 14, 24, 12, 60])
    pref = " > ".join(r["name"] for r in (meta.get("criterios") or {}).get("listados_referencia_prioridad") or []) or "Union Europea (UE) > PCPC > FDA"
    sw.title("Paso 2 · Composición", f"Ingredientes en nomenclatura INCI, en el orden del COCC. Listado de referencia: preferir {pref}; si la plataforma no ofrece el primero para un INCI, tomar el siguiente.")
    ingredients = (planilla.get("composition") or {}).get("ingredients") or []
    first = sw.row + 7
    last = first + max(len(ingredients), 1) - 1
    sw.section("Control de la fórmula", 13)
    r0 = sw.row
    ctrl = [
        ("Ingredientes registrados", f"=COUNTA($B${first}:$B${last})"),
        ("Suma ingredientes básicos", f'=SUMIFS($H${first}:$H${last},$D${first}:$D${last},"Básico")'),
        ("Suma secundarios / alérgenos", f'=SUMIFS($H${first}:$H${last},$D${first}:$D${last},"Secundario")'),
        ("Verificación (básicos = 100)", f'=IF(ABS(C{r0 + 1}-100)<0.005,"OK – suma 100","REVISAR – no suma 100")'),
    ]
    for k, f in ctrl:
        sw.ws.cell(sw.row, 1, k).font = _font(bold=True)
        c = sw.ws.cell(sw.row, 3, f)
        c.font = _font(bold=True)
        c.number_format = "0.000000"
        sw.row += 1
    sw.ws.cell(sw.row - 3, 5, "Los secundarios (alérgenos de la fragancia) normalmente ya están incluidos en el porcentaje de la fragancia.").font = _font(italic=True, color="595959")
    sw.blank()
    assert sw.row == first - 1, (sw.row, first)
    sw.header(cols)
    for ing in ingredients:
        fn = ing.get("functions") or []
        fnames = "; ".join(f.get("name") or "?" for f in fn)
        fslugs = "; ".join(f.get("slug") or "?" for f in fn)
        pend = bool(ing.get("_decision")) or any(f.get("slug") is None for f in fn) or (ing.get("referenceList") or {}).get("slug") is None or (ing.get("unit") or {}).get("slug") is None
        note = " ".join(x for x in [ing.get("_decision"), *(f.get("_decision") for f in fn), ing.get("_note")] if x)
        sw.line(
            [
                ing.get("order"),
                ing.get("inci"),
                ing.get("cas") or "",
                "Básico" if ing.get("itemType") == "basic" else "Secundario",
                fnames,
                fslugs,
                (ing.get("referenceList") or {}).get("name") or "PENDIENTE",
                ing.get("quantity"),
                (ing.get("unit") or {}).get("name") or "PENDIENTE",
                _val(ing.get("isNanomaterial")),
                ing.get("functionSource") or "",
                "Pendiente" if pend else "Listo",
                note,
            ],
            wrap_cols=(5, 13),
        )
        sw.ws.cell(sw.row - 1, 8).number_format = "0.000000"
        if pend:
            sw.ws.cell(sw.row - 1, 12).font = _font(bold=True, color=RED)
    ws.freeze_panes = ws.cell(first, 3)

    # --- 3 Información de uso ------------------------------------------------------
    ws = wb.create_sheet("3 Información de uso")
    sw = SheetWriter(ws, STEP_WIDTHS)
    use = planilla.get("useInformation") or {}
    sw.title("Paso 3 · Información de uso", use.get("_note") or "Textos tomados literalmente de la etiqueta o del sticker complementario.")
    sw.header(STEP_COLS)
    sw.section("3.1 Uso del producto")
    _field(sw, "3.1.1", "Modo de uso / instrucciones", use, "instructions")
    _field(sw, "3.1.2", "Advertencias", use, "warnings")

    # --- 4 Características técnicas --------------------------------------------------
    ws = wb.create_sheet("4 Caract. técnicas")
    sw = SheetWriter(ws, STEP_WIDTHS)
    tech = planilla.get("technicalCharacteristics") or {}
    sw.title("Paso 4 · Características técnicas", "Fuentes: especificaciones de producto terminado, declaración microbiológica, certificado de vida útil.")
    sw.header(STEP_COLS)
    sw.section("4.1 Características organolépticas")
    org = tech.get("organoleptic") or {}
    _field(sw, "4.1.1", "Color", org, "color")
    _field(sw, "4.1.2", "Olor", org, "odor")
    _field(sw, "4.1.3", "Sabor", org, "taste")
    others = tech.get("otherOrganoleptic") or []
    if others:
        sw.subsection("Otras características organolépticas")
        for i, o in enumerate(others, 1):
            sw.line([f"4.1.{3 + i}", f"Característica: {o.get('characteristic')}", o.get("value"), "Texto", "", "Listo", ""])
    sw.blank()
    sw.section("4.2 Características fisicoquímicas")
    for p in tech.get("physicochemical") or []:
        i = p.get("order")
        pend = bool(p.get("_decision")) or not p.get("parameterTypeSlug") or not p.get("unitSlug")
        sw.subsection(f"Parámetro {i} · {p.get('parameterTypeName')}")
        sw.line([f"4.2.{i}.1", "Tipo de parámetro", p.get("parameterTypeName"), "Lista", p.get("parameterTypeSlug") or "", "Pendiente" if not p.get("parameterTypeSlug") else "Listo", p.get("_note") or ""], pending=not p.get("parameterTypeSlug"))
        sw.line([f"4.2.{i}.2", "¿Es rango?", _val(p.get("isRange")), "Sí/No", "", "Listo", ""])
        if p.get("isRange"):
            sw.line([f"4.2.{i}.3", "Mínimo", p.get("min"), "Número", "", "Listo", ""])
            sw.line([f"4.2.{i}.4", "Máximo", p.get("max"), "Número", "", "Listo", ""])
        else:
            sw.line([f"4.2.{i}.3", "Valor", p.get("value"), "Número", "", "Listo", ""])
        sw.line([f"4.2.{i}.5", "Unidad", p.get("unitName") or "PENDIENTE", "Lista", p.get("unitSlug") or "", "Pendiente" if pend else "Listo", p.get("_decision") or ""], pending=pend)
    sw.blank()
    sw.section("4.3 Características microbiológicas")
    mic = tech.get("microbiological") or {}
    sw.line(["4.3.1", "¿Exceptuado de análisis microbiológico?", _val(mic.get("isExcepted")), "Sí/No", "", "Pendiente" if mic.get("isExcepted") is None else "Listo", mic.get("_decision") or mic.get("_note") or ""], pending=mic.get("isExcepted") is None or bool(mic.get("_decision")))
    if mic.get("justification"):
        sw.line(["4.3.2", "Justificación", mic.get("justification"), "Texto", "", "Listo", ""])
    _catfield(sw, "4.3.3", "Área de aplicación", mic.get("applicationArea"))
    sw.line(["4.3.4", "Fase etárea", mic.get("ageGroupName") or "PENDIENTE", "Lista", mic.get("ageGroupSlug") or "", "Listo" if mic.get("ageGroupSlug") else "Pendiente", ""], pending=not mic.get("ageGroupSlug"))
    sw.line(["4.3.5", "Límites declarados en especificaciones", mic.get("limits") or "", "Informativo", "", "Informativo", ""])
    sw.blank()
    sw.section("4.4 Vida útil")
    sl = tech.get("shelfLife") or {}
    sw.line(["4.4.1", "Vida útil (meses)", _val(sl.get("months")), "Número", "", "Pendiente" if sl.get("months") is None else "Listo", sl.get("_decision") or sl.get("_note") or ""], pending=sl.get("months") is None)

    # --- 5 Documentos obligatorios -----------------------------------------------------
    ws = wb.create_sheet("5 Doc. obligatorios")
    sw = SheetWriter(ws, [9, 40, 70, 12, 12, 16, 60])
    sw.title("Paso 5 · Documentos obligatorios", "Nombres de archivo copiados tal cual están en la carpeta. Un PDF de máximo 10 MB por casilla.")
    sw.header(["#", "Casilla del formulario", "Archivo a adjuntar (nombre literal)", "Tipo", "Tamaño (MB)", "Estado", "Nota / decisión"])
    md = planilla.get("mandatoryDocuments") or {}

    def files_block(num: str, label: str, files: list[str], note: str = "", pend: bool = False):
        sw.section(f"{num} {label}")
        if not files:
            sw.line([f"{num}.a", label, "PENDIENTE – no se encontró en la carpeta", "Adjunto", "", "Pendiente", note], pending=True)
            return
        multi = len(files) > 1
        for j, f in enumerate(files):
            letter = chr(ord("a") + j)
            lbl = label if j == 0 else "↳ misma casilla"
            st = "Pendiente" if (pend or multi) else "Listo"
            n = note if j == 0 else ("Varios archivos para una casilla: unir en un PDF" if multi else "")
            sw.line([f"{num}.{letter}", lbl, f, "Adjunto", _file_mb(folder, f), st, n], pending=(pend or multi))

    files_block("5.1", "Declaración de no realización de pruebas en animales", md.get("animalTestingDeclaration") or [])
    files_block("5.2", "Soporte de los beneficios cosméticos", md.get("cosmeticBenefitsSupport") or [])
    lab = md.get("label") or {}
    sw.section("5.3 Etiqueta")
    sw.line(["5.3.0", "¿Es la etiqueta definitiva?", _val(lab.get("isFinalLabel")), "Sí/No", "", "Pendiente" if lab.get("isFinalLabel") is None else "Listo", lab.get("_decision") or ""], pending=lab.get("isFinalLabel") is None)
    for j, f in enumerate(lab.get("artworkFiles") or []):
        sw.line([f"5.3.{chr(ord('a') + j)}", "Arte de etiqueta" if j == 0 else "↳ misma casilla", f, "Adjunto", _file_mb(folder, f), "Pendiente" if len(lab.get("artworkFiles") or []) > 1 else "Listo", "Varios archivos para una casilla: unir en un PDF < 10 MB" if j == 0 and len(lab.get("artworkFiles") or []) > 1 else ""], pending=len(lab.get("artworkFiles") or []) > 1)
    for k, lc in enumerate(md.get("lotsCoding") or [], 1):
        files_block(f"5.4.{k}", f"Codificación de lotes · {lc.get('manufacturer')}", lc.get("files") or [])
    ph = md.get("phosphorusContent") or {}
    files_block("5.5", "Contenido de fósforo · certificado analítico", ph.get("analyticalCertificate") or [])
    files_block("5.6", "Contenido de fósforo · declaración de conformidad", ph.get("conformityDeclaration") or [])
    if md.get("_note"):
        sw.line(["", "Nota", md["_note"], "", "", "Informativo", ""])

    # --- 6 Documentos adicionales -----------------------------------------------------
    ws = wb.create_sheet("6 Doc. adicionales")
    sw = SheetWriter(ws, [9, 46, 70, 12, 12, 16, 40])
    sw.title("Paso 6 · Documentos adicionales", "Un archivo por casilla salvo que se indique otra cosa. Confirmar la lista con el equipo.")
    sw.header(["#", "Casilla propuesta", "Archivo a adjuntar (nombre literal)", "Tipo", "Tamaño (MB)", "Estado", "Nota"])
    sw.section("6.1 Documentos adicionales propuestos")
    for j, d in enumerate(planilla.get("additionalDocuments") or [], 1):
        mb = _file_mb(folder, d.get("archivo", ""))
        big = mb not in ("no encontrado",) and float(mb) > 10
        sw.line([f"6.{j}", d.get("descripcion"), d.get("archivo"), "Adjunto", mb, "Pendiente" if big or mb == "no encontrado" else "Listo", "Supera 10 MB: comprimir" if big else ("Archivo no encontrado en la carpeta" if mb == "no encontrado" else "")], pending=big or mb == "no encontrado")

    # --- Decisiones ---------------------------------------------------------------------
    ws = wb.create_sheet("Decisiones")
    sw = SheetWriter(ws, [8, 28, 30, 70, 60, 60])
    sw.title("Decisiones pendientes", "Campos que no se pueden cerrar sin una decisión de la analista. Escriba la respuesta en la columna amarilla, guarde el Excel y use clic derecho → «Aplicar decisiones InvimÁgil».")
    sw.header(["ID", "Tema", "Dónde afecta", "Qué falta / hallazgo", "Propuesta", "RESPUESTA DE LA ANALISTA"])
    dec = planilla.get("decisiones") or []
    if not dec:
        sw.line(["—", "Sin decisiones pendientes", "", "", "", ""], wrap_cols=(4, 5, 6))
    for d in dec:
        sw.line([d.get("id"), d.get("tema"), d.get("donde"), d.get("hallazgo"), d.get("propuesta"), ""], wrap_cols=(2, 3, 4, 5, 6), input_col=6)
        sw.ws.cell(sw.row - 1, 1).font = _font(bold=True)
    if resueltas:
        sw.blank()
        sw.section("Decisiones ya resueltas (histórico)", 6)
        sw.header(["ID", "Tema", "Dónde afecta", "Hallazgo", "Propuesta", "Respuesta aplicada"])
        for r in resueltas:
            sw.line([r.get("id"), r.get("tema"), r.get("donde"), r.get("hallazgo"), r.get("propuesta"), r.get("respuesta")], wrap_cols=(2, 3, 4, 5, 6))

    # --- Archivos --------------------------------------------------------------------------
    ws = wb.create_sheet("Archivos")
    sw = SheetWriter(ws, [70, 22, 10, 80])
    sw.title("Archivos de la carpeta", "Cómo clasificó la herramienta cada archivo. Verifique que no falte ninguno.")
    sw.header(["Archivo", "Clasificación", "Legible", "Resumen"])
    for a in planilla.get("archivos") or []:
        sw.line([a.get("archivo"), a.get("clasificacion"), _val(a.get("legible")), a.get("resumen")], wrap_cols=(1, 4), pending=not a.get("legible"))
        if not a.get("legible"):
            sw.ws.cell(sw.row - 1, 3).font = _font(bold=True, color=RED)

    # --- Listas ------------------------------------------------------------------------------
    ws = wb.create_sheet("Listas")
    sw = SheetWriter(ws, [14, 46, 3, 14, 24, 3, 10, 24])
    sw.header(["Código función", "Función del ingrediente", "", "Código", "Listado de referencia", "", "Código", "Unidad de composición"])
    fun = sorted(cat.entries("cosmFun").values(), key=lambda e: e.name)
    ref = list(cat.entries("refList").values())
    uni = list(cat.entries("compUni").values())
    for i in range(max(len(fun), len(ref), len(uni))):
        if i < len(fun):
            sw.ws.cell(sw.row, 1, fun[i].slug).font = _font()
            sw.ws.cell(sw.row, 2, fun[i].name).font = _font()
        if i < len(ref):
            sw.ws.cell(sw.row, 4, ref[i].slug).font = _font()
            sw.ws.cell(sw.row, 5, ref[i].name).font = _font()
        if i < len(uni):
            sw.ws.cell(sw.row, 7, uni[i].slug).font = _font()
            sw.ws.cell(sw.row, 8, uni[i].name).font = _font()
        sw.row += 1
    # dropdown on ingredient type
    dv = DataValidation(type="list", formula1='"Básico,Secundario"', allow_blank=True)
    wb["2 Composición"].add_data_validation(dv)
    dv.add(f"D{first}:D{last}")

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def read_answers(xlsx_path: Path) -> list[dict]:
    """Read the Decisiones sheet: rows with a non-empty answer in column F (pending block only)."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, data_only=True)
    if "Decisiones" not in wb.sheetnames:
        return []
    ws = wb["Decisiones"]
    out = []
    for row in ws.iter_rows(min_row=1, values_only=True):
        if not row or row[0] in (None, "ID", "—"):
            continue
        if isinstance(row[0], str) and row[0].startswith("Decisiones ya resueltas"):
            break
        if len(row) >= 6 and row[5] not in (None, ""):
            out.append({"id": str(row[0]), "tema": row[1] or "", "donde": row[2] or "", "hallazgo": row[3] or "", "propuesta": row[4] or "", "respuesta": str(row[5]).strip()})
    return out
